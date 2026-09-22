import tempfile,json,itertools
from pathlib import Path
from datetime import date,timedelta
from unittest.mock import patch
from PySide6.QtWidgets import QApplication,QMessageBox,QFileDialog
from PySide6.QtCore import Qt
from PySide6.QtGui import QAction
from openpyxl import load_workbook
import taskmanager.db as db
from taskmanager.ui import MainWindow
from taskmanager.app import configure_main_window
from taskmanager import style_v8
app=QApplication([]);results=[];windows=[]
def run(repeat):
 root=Path(tempfile.mkdtemp());db.APP_DIR=root;db.DB_PATH=root/'db';db.BACKUP_DIR=root/'backups';db.init_db()
 w=configure_main_window(MainWindow());windows.append(w);w.resize(1440,900);w.show();app.processEvents()
 def check(name,ok,detail=''):results.append(dict(run=repeat,test=name,passed=bool(ok),detail=detail))
 def seed(title='Basis',due=None,status='Offen',priority='important_urgent',pid=None):return db.save(dict(title=title,description='Beschreibung',project_id=pid,priority=priority,due_date=due,status=status,recurrence='weekly',subtasks=[('Teil',True)]),make_backup=False)
 def reset():
  w.selected_task=None;w.editor_task_id=None;w._v8_detail_panel.clear_task();c=db.connect();c.execute('DELETE FROM tasks');c.execute('DELETE FROM projects');c.commit();c.close()
  for p in db.BACKUP_DIR.glob('*.db'):p.unlink()
  w._undo_available=False;w.search.clear();w.set_project(None);w.pfilter.setCurrentIndex(0);w.set_scope('Alle');style_v8._set_group(w,'Keine Gruppierung');w.set_view('tasks');w.refresh_all();app.processEvents()
 try:
  for operation in ['save','priority','duplicate','status','delete','planning']:
   reset();tid=seed();w.refresh_all();w.select_task(tid)
   if operation=='save':w.e_title.setText('Neu');w.save_editor()
   elif operation=='priority':w.cycle_priority(tid)
   elif operation=='duplicate':w.duplicate_selected()
   elif operation=='status':w.set_status(tid,'Erledigt')
   elif operation=='planning':w.drop_moved(tid,'Später')
   else:
    with patch.object(QMessageBox,'question',return_value=QMessageBox.Yes):w.delete_selected()
   w.undo_button.click();r=db.task(tid)
   check('undo_'+operation,r is not None and r['title']=='Basis' and r['priority']=='important_urgent' and r['status']=='Offen' and r['due_date'] is None and len(db.tasks())==1)
  reset();tid=seed();w.refresh_all();w.select_task(tid);w.e_title.setText('Neu');w.save_editor();check('save_persistence',db.task(tid)['title']=='Neu' and db.subs(tid)[0]['done']==1)
  w.e_title.setText('Nicht gespeichert');w.search.setText('Neu');check('unsaved_editor_survives_search',w.e_title.text()=='Nicht gespeichert')
  reset();tid=seed();w.refresh_all();w.select_task(tid)
  w.priority_cards['important_not_urgent'].click();w.select_task(tid);check('cancel_priority_restores_original',db.task(tid)['priority']=='important_urgent')
  reset();tid=seed(due='2027-01-01');w.refresh_all();w.set_status(tid,'Erledigt');check('recurring_task_has_next_occurrence',any(x['status']!='Erledigt' and x['due_date']=='2027-01-08' for x in db.tasks()))
  reset();tid=seed();w.refresh_all()
  for view,lanes in [('kanban',w.kcols),('eisenhower',w.ecols),('planning',w.pcols)]:
   style_v8._navigate(w,view);lane=next(x for x in lanes.values() if x.count());w._v8_detail_panel.set_panel_open(False);lane.itemClicked.emit(lane.item(0));app.processEvents();check('visible_inspector_'+view,w.editor.isVisible());check('selected_task_'+view,w.editor_task_id==tid)
  reset();early=seed('Januar',due='2027-01-31');late=seed('Februar',due='2027-02-01');w.refresh_all()
  for mode in [0,1,2]:
   style_v8._set_sort_mode(w,mode);ids=[w.table.item(i,1).data(Qt.UserRole) for i in range(w.table.rowCount())];check('sort_'+str(mode),ids==([late,early] if mode==2 else [early,late]))
  style_v8._set_group(w,'Fälligkeit');check('chronological_grouping',[w.table.item(i,1).data(Qt.UserRole) for i in range(2)]==[early,late])
  reset();today=date.today();end=today+timedelta(days=6-today.weekday());fixture=[]
  for priority,status,due in itertools.product(['important_urgent','important_not_urgent'],['Offen','Erledigt'],[None,today.isoformat(),(end+timedelta(days=1)).isoformat()]):
   tid=seed(status=status,priority=priority,due=due);fixture.append((tid,priority,status,due))
  w.refresh_all();good=True;n=0
  for pi,si,di in itertools.product(range(w.pfilter.count()),range(w.sfilter.count()),range(w.dfilter.count())):
   w.pfilter.setCurrentIndex(pi);w.sfilter.setCurrentIndex(si);w.dfilter.setCurrentIndex(di)
   expected={tid for tid,p,s,d in fixture if (pi==0 or p==w.pfilter.currentData()) and (si==0 or s==w.sfilter.currentText()) and (di==0 or di==1 and d==today.isoformat() or di==2 and d is not None and today.isoformat()<=d<=end.isoformat() or di==3 and d is not None and d>end.isoformat() or di==4 and d is None)}
   actual={w.table.item(i,1).data(Qt.UserRole) for i in range(w.table.rowCount())};good &= actual==expected;n+=1
  check('filter_combinations',good,str(n))
  w.pfilter.setCurrentIndex(0);w.dfilter.setCurrentIndex(0);w.sfilter.setCurrentText('Erledigt');check('board_respects_active_status_filter',all(w.kcols[s].count()==0 for s in ['Offen','In Arbeit']))
  reset();tid=seed();w.refresh_all();w.set_scope('Erledigt');w.sfilter.setCurrentIndex(0);check('scope_matches_cleared_status',w.scope!='Erledigt')
  reset();seed();w.refresh_all();before=w.table.rowHeight(0);style_v8._toggle_compact(w);check('compact_changes_existing_row_height',w.table.rowHeight(0)!=before)
  with patch.object(QFileDialog,'getSaveFileName',return_value=(str(root/'export.xlsx'),'')),patch.object(QMessageBox,'information'):
   w.export_excel()
  wb=load_workbook(root/'export.xlsx');check('excel_export',wb.active.max_row==2);wb.close()
  check('database_integrity',db.connect().execute('PRAGMA integrity_check').fetchone()[0]=='ok')
 finally:w.close()
for i in range(3):run(i+1)
print(json.dumps(results,ensure_ascii=False,indent=2));print('SUMMARY',len(results),sum(x['passed'] for x in results))
