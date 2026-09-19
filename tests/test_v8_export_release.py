import importlib.util
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

import taskmanager.db as db
import taskmanager.ui as ui


def test_clean_release_exports_real_workbook(tmp_path, monkeypatch):
    assert importlib.util.find_spec("openpyxl") is not None, (
        "openpyxl must be installed by the release requirements"
    )
    from openpyxl import load_workbook

    monkeypatch.setattr(db, "APP_DIR", tmp_path)
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "tasks.db")
    monkeypatch.setattr(db, "BACKUP_DIR", tmp_path / "backups")
    db.init_db()
    db.save(
        {
            "title": "Exportprüfung",
            "description": "Releaseinhalt",
            "project_id": None,
            "priority": "important_urgent",
            "due_date": "2026-09-19",
            "status": "Offen",
            "recurrence": "none",
            "subtasks": [("Prüfschritt", True)],
        },
        make_backup=False,
    )

    target = tmp_path / "Arturs_Taskmanager_Aufgaben.xlsx"
    app = QApplication.instance() or QApplication([])
    window = ui.MainWindow()
    monkeypatch.setattr(
        QFileDialog,
        "getSaveFileName",
        staticmethod(lambda *args, **kwargs: (str(target), "Excel-Dateien (*.xlsx)")),
    )
    monkeypatch.setattr(
        QMessageBox,
        "information",
        staticmethod(lambda *args, **kwargs: QMessageBox.Ok),
    )

    try:
        window.export_excel()
    finally:
        window.close()
        window.deleteLater()
        app.processEvents()

    assert target.exists()
    workbook = load_workbook(target, read_only=True)
    rows = list(workbook["Aufgaben"].iter_rows(values_only=True))
    assert rows[1][0] == "Exportprüfung"
    assert rows[1][5] == "Releaseinhalt"
    assert rows[1][6] == "1/1"
